"""
RATE_DXY.xlsx 생성 — 시장금리(1016_rate.xlsx)와 달러인덱스(Investing.com 일별 CSV 2개) 전처리

입력 (같은 폴더)
  1016_rate.xlsx                          NECOS 721U001 시장금리 월·분기·년 (2열씩 날짜·값, 13행부터 자료)
  US Dollar Index Historical Data (1).csv  Investing.com DXY 일별 1995-01-02 ~ 2014-05-14
  US Dollar Index Historical Data (2).csv  Investing.com DXY 일별 2014-05-14 ~
출력
  RATE_DXY.xlsx  시트 월별 / DXY_일별 / META
  build_DATA.py가 이 파일을 읽는다

필요: pandas, numpy, openpyxl
"""

import os
import glob
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

HERE = os.path.dirname(os.path.abspath(__file__)) if "__file__" in globals() else os.getcwd()
OUT  = os.path.join(HERE, "RATE_DXY.xlsx")


# 1. 시장금리: 월 계열만 쓴다 (ECOS 월 통계 = 월평균)

RATE_CODES = {
    "5020000": "ktb3y",   # 국고채(3년)
    "6010100": "msb1y",   # 통안증권(1년)
    "7020000": "aa3y",    # 회사채(3년, AA-)
    "7030000": "bbb3y",   # 회사채(3년, BBB-), 2000M10부터
}

raw = pd.read_excel(os.path.join(HERE, "1016_rate.xlsx"), header=None)
rate = {}
for c in range(0, raw.shape[1] - 1, 2):
    code = str(raw.iloc[0, c])
    parts = code.split("-")
    if len(parts) < 4 or parts[2] != "M" or parts[3] not in RATE_CODES:
        continue
    d = raw.iloc[13:, [c, c + 1]].copy()
    d.columns = ["d", "v"]
    d["v"] = pd.to_numeric(d["v"], errors="coerce")
    d = d.dropna()
    d["d"] = pd.to_datetime(d["d"])
    rate[RATE_CODES[parts[3]]] = d.set_index("d")["v"].sort_index()

miss = [v for v in RATE_CODES.values() if v not in rate]
if miss:
    raise SystemExit(f"1016_rate.xlsx에 월 계열이 없다: {miss}")

R = pd.DataFrame(rate)
R["spread"]     = R["aa3y"] - R["ktb3y"]     # 메인: AA- − 국고 (기존 SPREAD_LONG과 같은 정의)
R["spread_bbb"] = R["bbb3y"] - R["aa3y"]     # 강건성: BBB- − AA-


# 2. 달러인덱스: 두 CSV를 이어 붙이고 일별 → 월평균

files = sorted(glob.glob(os.path.join(HERE, "US Dollar Index Historical Data*.csv")))
if not files:
    raise SystemExit("US Dollar Index Historical Data*.csv가 없다")

parts = []
for f in files:
    x = pd.read_csv(f, encoding="utf-8-sig")
    x["Date"]  = pd.to_datetime(x["Date"], format="%m/%d/%Y")
    x["Price"] = pd.to_numeric(x["Price"].astype(str).str.replace(",", ""), errors="coerce")
    parts.append(x[["Date", "Price"]].dropna())
dxy_all = pd.concat(parts)

# 두 파일이 겹치는 날은 값이 같아야 한다
dup = dxy_all[dxy_all["Date"].duplicated(keep=False)]
if not dup.empty and dup.groupby("Date")["Price"].nunique().max() > 1:
    raise SystemExit(f"겹치는 날짜의 값이 다르다:\n{dup.sort_values('Date')}")

dxy_d = dxy_all.drop_duplicates("Date").sort_values("Date").set_index("Date")["Price"]
dxy_m = dxy_d.resample("MS").agg(["mean", "last", "count"])
dxy_m.columns = ["dxy", "dxy_eop", "dxy_ndays"]

# 마지막 달은 불완전 월일 수 있다. 다운로드일이 속한 달은 버린다
last_month = dxy_d.index.max().to_period("M").to_timestamp()
dxy_m = dxy_m[dxy_m.index < last_month]


# 3. 합치기

M = R.join(dxy_m, how="outer").sort_index()
M = M.loc["1995-01-01":]
M.insert(0, "ym", [f"{d.year}m{d.month}" for d in M.index])
M.index.name = "date"

meta = pd.DataFrame([
    ("ktb3y",      "국고채 3년 (%)",                    "1016_rate.xlsx", "NECOS 721U001 M 5020000", "월"),
    ("msb1y",      "통안증권 1년 (%)",                  "1016_rate.xlsx", "NECOS 721U001 M 6010100", "월"),
    ("aa3y",       "회사채 3년 AA- (%)",                "1016_rate.xlsx", "NECOS 721U001 M 7020000", "월"),
    ("bbb3y",      "회사채 3년 BBB- (%)",               "1016_rate.xlsx", "NECOS 721U001 M 7030000", "월, 2000M10부터"),
    ("spread",     "AA- − 국고 3년 (%p)",               "계산",           "aa3y − ktb3y",            "월"),
    ("spread_bbb", "BBB- − AA- 3년 (%p)",               "계산",           "bbb3y − aa3y",            "월, 2000M10부터"),
    ("dxy",        "ICE 달러인덱스 월평균",             "Investing.com", "US Dollar Index (DXY)",   "일별 → 월평균"),
    ("dxy_eop",    "ICE 달러인덱스 월말",               "Investing.com", "US Dollar Index (DXY)",   "일별 → 월말"),
    ("dxy_ndays",  "달러인덱스 월별 관측일수",          "Investing.com", "—",                        "—"),
], columns=["변수", "설명", "원천", "계열", "월 변환"])


# 4. 저장

with pd.ExcelWriter(OUT, engine="openpyxl") as w:
    M.to_excel(w, sheet_name="월별")
    dxy_d.rename("dxy").to_frame().to_excel(w, sheet_name="DXY_일별")
    meta.to_excel(w, sheet_name="META", index=False)

wb = load_workbook(OUT)
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            c.font = Font(name="Arial", size=10, bold=(c.row == 1))
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 12
for c in wb["월별"]["A"][1:]:
    c.number_format = "yyyy-mm"
for c in wb["DXY_일별"]["A"][1:]:
    c.number_format = "yyyy-mm-dd"
wb.save(OUT)

print(f"저장: {OUT}")
print(f"  시장금리 {R.index.min():%Y-%m} ~ {R.index.max():%Y-%m}, BBB- 시작 {R['bbb3y'].first_valid_index():%Y-%m}")
print(f"  DXY 일별 {dxy_d.index.min():%Y-%m-%d} ~ {dxy_d.index.max():%Y-%m-%d} ({len(dxy_d)}일), 월 {dxy_m.index.min():%Y-%m} ~ {dxy_m.index.max():%Y-%m}")
